import shutil
import subprocess
import time
import math
from pathlib import Path

KPI_TOLERANCE_PERCENT = 1.0
DISPLAY_DECIMAL_PLACES = 3
EXPECTED_DEBUG_COLUMNS = [
    "code",
    "pdf",
    "cdf",
    "cos_cdf",
    "dnl",
    "inl",
    "inl_polynomial",
]
EXPECTED_KPIS = {
    "inl_codes_p2p": 0.1843,
    "inl_max": 3.307,
    "code_inl_max": 1171,
    "inl_min": -2.4758,
    "code_inl_min": 874,
    "missing_codes_threshold": -0.9,
    "missing_codes_count": 0,
    "dnl_max": 1.0015,
    "code_dnl_max": 957,
    "dnl_min": -0.8174,
    "code_dnl_min": 982,
    "dnl_rms": 0.1614,
    "code_min": 779,
    "code_max": 1273,
    "code_trunclow": 787,
    "code_trunchigh": 1264,
    "lsb_codes_over_code_amp": 0.004137,
}


def add_check(report, name, passed, reason=None, skipped=False):
    item = {"check": name, "passed": bool(passed), "skipped": bool(skipped)}
    if reason:
        item["reason"] = reason
    report.append(item)


def check_status(report):
    failed = sum(1 for x in report if (not x.get("skipped", False) and not x.get("passed", False)))
    passed = sum(1 for x in report if (not x.get("skipped", False) and x.get("passed", False)))
    skipped = sum(1 for x in report if x.get("skipped", False))
    return passed, failed, skipped


def pct_diff(golden, computed):
    if golden == 0:
        return 0.0 if computed == 0 else None
    return abs(computed - golden) / abs(golden) * 100.0


def find_node():
    node = shutil.which("node")
    if node:
        return node

    candidates = [
        Path("/usr/bin/node"),
        Path("/usr/local/bin/node"),
    ]
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return str(candidate)

    nvm_root = Path.home() / ".nvm" / "versions" / "node"
    if nvm_root.exists():
        versions = list(nvm_root.glob("*/bin/node"))
        if versions:
            versions.sort(key=lambda p: str(p), reverse=True)
            return str(versions[0])

    raise RuntimeError("Node.js was not found. Please install Node.js before running this test.")


def run_command(command):
    command = [str(x) for x in command]
    try:
        return subprocess.run(command, capture_output=True, text=True)
    except Exception as exc:
        return exc


def parse_scalar_outputs(stdout_text):
    outputs = {}
    if not stdout_text:
        return outputs
    in_outputs = False
    for raw in stdout_text.splitlines():
        line = raw.strip()
        if line == "Outputs:":
            in_outputs = True
            continue
        if not in_outputs or not line or ":" not in line:
            continue
        key, value = line.split(":", 1)
        key, value = key.strip(), value.strip()
        if not key:
            continue
        try:
            outputs[key] = float(value)
        except ValueError:
            outputs[key] = value
    return outputs


def save_text(path, text):
    path.write_text(text or "", encoding="utf-8")


def pd_to_docx_table(document, dataframe):
    table = document.add_table(rows=dataframe.shape[0] + 1, cols=dataframe.shape[1])
    table.style = "Table Grid"
    for col_idx, col in enumerate(dataframe.columns):
        table.cell(0, col_idx).text = str(col)
    for row_idx in range(dataframe.shape[0]):
        for col_idx in range(dataframe.shape[1]):
            value = dataframe.iat[row_idx, col_idx]
            table.cell(row_idx + 1, col_idx).text = "" if value is None else str(value)
    return table


def format_num(value):
    try:
        if value is None:
            return "N/A"
        numeric = float(value)
        if numeric != numeric:
            return "NaN"
        rounded = round(numeric, DISPLAY_DECIMAL_PLACES)
        if rounded == 0:
            return "0"
        if float(rounded).is_integer():
            return str(int(rounded))
        text = f"{rounded:.{DISPLAY_DECIMAL_PLACES}f}".rstrip("0").rstrip(".")
        return text if text != "-0" else "0"
    except Exception:
        return str(value)


def main():
    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError("pandas is required for this test.") from exc

    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError("python-docx is required for this test.") from exc

    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        plt = None

    project_root = Path(__file__).resolve().parents[2]
    results_root = project_root / "test" / "results" / "sinl_tests_docx"
    timestamp = time.strftime("%Y_%m_%d_%H_%M_%S")
    output_folder = results_root / timestamp
    output_folder.mkdir(parents=True, exist_ok=True)

    report = []
    test_rows = []

    def add_case(name, passed, reason=None, skipped=False):
        add_check(report, name, passed, reason=reason, skipped=skipped)
        test_rows.append({"Test": name, "Status": "SKIP" if skipped else ("PASS" if passed else "FAIL"), "Reason": reason or ""})

    usig_cli = project_root / "usig.mjs"
    golden_raw_csv = project_root / "test" / "golden_raw_data" / "sine_fs2p25ghz_tonemode~single_fftlength8192_numaveraging4_numberofcores8_ticorrections~ogp.csv"
    golden_debug_csv = project_root / "test" / "golden_raw_data" / "inl_dnl_series.csv"
    computed_debug_csv = output_folder / "inl_dnl_series_computed.csv"

    node_executable = None
    node_version = "N/A"
    cli_result = None

    try:
        node_executable = find_node()
        add_case("Node.js discovery", True)
        nv = subprocess.run([node_executable, "--version"], capture_output=True, text=True)
        node_version = (nv.stdout or "").strip() if nv.returncode == 0 else "unknown"
        add_case("Node.js version query", nv.returncode == 0, None if nv.returncode == 0 else f"exit code {nv.returncode}")
    except Exception as exc:
        add_case("Node.js discovery", False, str(exc))

    add_case("USIG CLI path", usig_cli.exists() and usig_cli.is_file(), None if usig_cli.exists() else f"missing {usig_cli}")
    add_case("Golden raw dataset", golden_raw_csv.exists() and golden_raw_csv.is_file(), None if golden_raw_csv.exists() else f"missing {golden_raw_csv}")
    add_case("Golden debug dataset", golden_debug_csv.exists() and golden_debug_csv.is_file(), None if golden_debug_csv.exists() else f"missing {golden_debug_csv}")

    scalar_outputs = {}
    scalar_result = None
    scalar_command = None
    if node_executable and usig_cli.exists() and golden_raw_csv.exists():
        scalar_command = [node_executable, str(usig_cli), "-i", str(golden_raw_csv), "-plugin", "sinl"]
        scalar_result = run_command(scalar_command)
        if isinstance(scalar_result, Exception):
            add_case("CLI scalar KPI execution", False, f"failed to start: {scalar_result}")
            save_text(output_folder / "scalar_command_stdout.txt", "")
            save_text(output_folder / "scalar_command_stderr.txt", str(scalar_result))
        else:
            add_case("CLI scalar KPI execution", scalar_result.returncode == 0, None if scalar_result.returncode == 0 else f"exit code {scalar_result.returncode}")
            save_text(output_folder / "scalar_command_stdout.txt", scalar_result.stdout)
            save_text(output_folder / "scalar_command_stderr.txt", scalar_result.stderr)
            scalar_outputs = parse_scalar_outputs(scalar_result.stdout)
            add_case("Scalar Outputs section", len(scalar_outputs) > 0, None if scalar_outputs else "Outputs section not found")
    else:
        add_case("CLI scalar KPI execution", True, "prerequisite missing", skipped=True)
        add_case("Scalar Outputs section", True, "scalar command skipped", skipped=True)

    kpi_rows = []
    for kpi, golden_value in EXPECTED_KPIS.items():
        actual = scalar_outputs.get(kpi)
        if actual is None:
            kpi_rows.append({"KPI": kpi, "Golden": format_num(golden_value), "Computed": "N/A", "Difference %": "N/A", "Status": "FAIL"})
            add_case(f"KPI {kpi}", False, "missing in scalar outputs")
            continue
        try:
            actual_f = float(actual)
            golden_f = float(golden_value)
            diff = pct_diff(golden_f, actual_f)
            if diff is None:
                passed = False
                diff_str = "undefined (golden=0, computed!=0)"
            else:
                passed = diff <= KPI_TOLERANCE_PERCENT
                diff_str = format_num(diff)

            kpi_rows.append({
                "KPI": kpi,
                "Golden": format_num(golden_f),
                "Computed": format_num(actual_f),
                "Difference %": diff_str,
                "Status": "PASS" if passed else "FAIL",
            })

            add_case(
                f"KPI {kpi}",
                passed,
                None if passed else f"difference {diff_str}% > {KPI_TOLERANCE_PERCENT}%"
            )
        except Exception as exc:
            kpi_rows.append({
                "KPI": kpi,
                "Golden": format_num(golden_value),
                "Computed": format_num(actual),
                "Difference %": "N/A",
                "Status": "FAIL",
            })
            add_case(f"KPI {kpi}", False, f"non-numeric: {exc}")

    cli_command = None
    if node_executable and usig_cli.exists() and golden_raw_csv.exists():
        cli_command = [
            node_executable,
            str(usig_cli),
            "-i",
            str(golden_raw_csv),
            "-plugin",
            "sinl",
            "-debug",
            f"inl_dnl_series={computed_debug_csv}",
        ]
        cli_result = run_command(cli_command)
        if isinstance(cli_result, Exception):
            add_case("CLI debug table execution", False, f"failed to start: {cli_result}")
            save_text(output_folder / "command_stdout.txt", "")
            save_text(output_folder / "command_stderr.txt", str(cli_result))
        else:
            add_case("CLI debug table execution", cli_result.returncode == 0, None if cli_result.returncode == 0 else f"exit code {cli_result.returncode}")
            save_text(output_folder / "command_stdout.txt", cli_result.stdout)
            save_text(output_folder / "command_stderr.txt", cli_result.stderr)
    else:
        add_case("CLI debug table execution", True, "prerequisite missing", skipped=True)
        save_text(output_folder / "command_stdout.txt", "")
        save_text(output_folder / "command_stderr.txt", "CLI command skipped due to missing prerequisites")

    computed_exists = computed_debug_csv.exists() and computed_debug_csv.is_file()
    add_case("Debug table generated", computed_exists, None if computed_exists else f"missing {computed_debug_csv}")

    golden_df = None
    computed_df = None

    if golden_debug_csv.exists():
        try:
            golden_df = pd.read_csv(golden_debug_csv)
            add_case("Golden debug table load", True)
        except Exception as exc:
            add_case("Golden debug table load", False, str(exc))
    else:
        add_case("Golden debug table load", False, "golden debug file missing")

    if computed_exists:
        try:
            computed_df = pd.read_csv(computed_debug_csv)
            add_case("Computed debug table load", True)
        except Exception as exc:
            add_case("Computed debug table load", False, str(exc))
    else:
        add_case("Computed debug table load", True, "computed debug file missing", skipped=True)

    structure_ok = True
    if golden_df is None or computed_df is None:
        structure_ok = False
        add_case("Required columns", True, "table load failed", skipped=True)
        add_case("Row count", True, "table load failed", skipped=True)
    else:
        missing_golden_cols = [c for c in EXPECTED_DEBUG_COLUMNS if c not in golden_df.columns]
        missing_computed_cols = [c for c in EXPECTED_DEBUG_COLUMNS if c not in computed_df.columns]
        has_columns = (len(missing_golden_cols) == 0 and len(missing_computed_cols) == 0)
        structure_ok = structure_ok and has_columns
        reason = None
        if not has_columns:
            reason = f"missing golden={missing_golden_cols}, missing computed={missing_computed_cols}"
        add_case("Required columns", has_columns, reason)

        same_rows = len(golden_df.index) == len(computed_df.index)
        structure_ok = structure_ok and same_rows
        add_case("Row count", same_rows, None if same_rows else f"golden={len(golden_df.index)}, computed={len(computed_df.index)}")

    vector_stats_rows = []
    plot_files = {}

    for vector in EXPECTED_DEBUG_COLUMNS:
        if not structure_ok:
            add_case(f"{vector} statistics", True, "skipped due to structural failure", skipped=True)
            vector_stats_rows.append({
                "Vector": vector,
                "Golden Mean": "N/A",
                "Computed Mean": "N/A",
                "Mean Diff %": "N/A",
                "Golden Std": "N/A",
                "Computed Std": "N/A",
                "Std Diff %": "N/A",
                "Mean": "SKIP",
                "Std": "SKIP",
                "Overall": "SKIP",
                "Interpretation": "Skipped due to structural validation failure.",
                "Golden NaNs": "N/A",
                "Computed NaNs": "N/A",
            })
            continue

        g = golden_df[vector]
        c = computed_df[vector]

        g_nan = int(g.isna().sum())
        c_nan = int(c.isna().sum())
        nan_count_match = g_nan == c_nan
        nan_location_match = bool(g.isna().equals(c.isna()))

        g_valid = g.dropna()
        c_valid = c.dropna()
        if g_valid.empty or c_valid.empty:
            add_case(
                f"{vector} statistics",
                False,
                "no valid numeric values after NaN removal",
            )
            vector_stats_rows.append({
                "Vector": vector,
                "Golden Mean": "N/A",
                "Computed Mean": "N/A",
                "Mean Diff %": "N/A",
                "Golden Std": "N/A",
                "Computed Std": "N/A",
                "Std Diff %": "N/A",
                "Mean": "FAIL",
                "Std": "FAIL",
                "Overall": "FAIL",
                "Interpretation": "Failed: no valid numeric values after ignoring NaNs.",
                "Golden NaNs": g_nan,
                "Computed NaNs": c_nan,
                "NaN Count Match": "YES" if nan_count_match else "NO",
                "NaN Location Match": "YES" if nan_location_match else "NO",
            })
            continue

        g_mean = float(g.mean(skipna=True))
        c_mean = float(c.mean(skipna=True))
        g_std = float(g.std(skipna=True, ddof=1))
        c_std = float(c.std(skipna=True, ddof=1))

        if any(x != x for x in [g_mean, c_mean, g_std, c_std]):
            add_case(
                f"{vector} statistics",
                False,
                "mean/std is NaN after NaN-aware calculation",
            )
            vector_stats_rows.append({
                "Vector": vector,
                "Golden Mean": "NaN",
                "Computed Mean": "NaN",
                "Mean Diff %": "N/A",
                "Golden Std": "NaN",
                "Computed Std": "NaN",
                "Std Diff %": "N/A",
                "Mean": "FAIL",
                "Std": "FAIL",
                "Overall": "FAIL",
                "Interpretation": "Failed: insufficient valid data for mean/std comparison.",
                "Golden NaNs": g_nan,
                "Computed NaNs": c_nan,
                "NaN Count Match": "YES" if nan_count_match else "NO",
                "NaN Location Match": "YES" if nan_location_match else "NO",
            })
            continue

        mean_diff = c_mean - g_mean
        std_diff = c_std - g_std

        mean_pct = pct_diff(g_mean, c_mean)
        std_pct = pct_diff(g_std, c_std)

        mean_pass = (mean_pct is not None and mean_pct <= KPI_TOLERANCE_PERCENT)
        std_pass = (std_pct is not None and std_pct <= KPI_TOLERANCE_PERCENT)
        overall_pass = mean_pass and std_pass

        add_case(f"{vector} statistics", overall_pass, None if overall_pass else f"mean_diff_pct={mean_pct}, std_diff_pct={std_pct}")

        interpretation = (
            f"Both mean and standard deviation are within {KPI_TOLERANCE_PERCENT}% tolerance."
            if overall_pass else
            f"At least one metric is outside {KPI_TOLERANCE_PERCENT}% tolerance."
        )

        vector_stats_rows.append({
            "Vector": vector,
            "Golden Mean": format_num(g_mean),
            "Computed Mean": format_num(c_mean),
            "Mean Diff %": "undefined" if mean_pct is None else format_num(mean_pct),
            "Golden Std": format_num(g_std),
            "Computed Std": format_num(c_std),
            "Std Diff %": "undefined" if std_pct is None else format_num(std_pct),
            "Mean": "PASS" if mean_pass else "FAIL",
            "Std": "PASS" if std_pass else "FAIL",
            "Overall": "PASS" if overall_pass else "FAIL",
            "Interpretation": interpretation,
            "Golden NaNs": g_nan,
            "Computed NaNs": c_nan,
            "NaN Count Match": "YES" if nan_count_match else "NO",
            "NaN Location Match": "YES" if nan_location_match else "NO",
            "golden_mean_raw": g_mean,
            "computed_mean_raw": c_mean,
            "mean_diff_raw": mean_diff,
            "golden_std_raw": g_std,
            "computed_std_raw": c_std,
            "std_diff_raw": std_diff,
            "mean_pct_raw": mean_pct,
            "std_pct_raw": std_pct,
        })

        if plt is not None:
            diff = c - g
            diff_clean = diff.dropna()
            diff_x = diff_clean.index.to_numpy()
            diff_values = diff_clean.to_numpy()
            max_abs = float(max(abs(diff_values.min()), abs(diff_values.max()))) if len(diff_values) > 0 else 0.0
            if not math.isfinite(max_abs) or max_abs == 0:
                ylim = 1.0
            else:
                ylim = max_abs * 1.1

            fig, ax = plt.subplots(figsize=(10, 4.2))
            ax.plot(diff_x, diff_values, label="computed - golden", linewidth=1.2)
            ax.axhline(0.0, color="red", linestyle="--", linewidth=1.0, label="y = 0")
            ax.set_title(f"SINL Debug Vector Difference: {vector}")
            ax.set_xlabel("vector index")
            ax.set_ylabel("computed - golden")
            ax.set_ylim(-ylim, ylim)
            ax.grid(True, alpha=0.3)
            ax.legend(loc="best")
            diff_plot_path = output_folder / f"{vector}_difference.png"
            fig.tight_layout()
            fig.savefig(diff_plot_path, dpi=160)
            plt.close(fig)

            fig2, ax2 = plt.subplots(figsize=(10, 4.2))
            ax2.plot(g.index.to_numpy(), g.to_numpy(), label="golden", linewidth=1.1, linestyle="-")
            ax2.plot(c.index.to_numpy(), c.to_numpy(), label="computed", linewidth=1.1, linestyle="--", alpha=0.85)
            ax2.set_title(f"SINL Debug Vector Comparison: {vector}")
            ax2.set_xlabel("vector index")
            ax2.set_ylabel(vector)
            ax2.grid(True, alpha=0.3)
            ax2.legend(loc="best")
            cmp_plot_path = output_folder / f"{vector}_comparison.png"
            fig2.tight_layout()
            fig2.savefig(cmp_plot_path, dpi=160)
            plt.close(fig2)

            plot_files[vector] = {
                "difference": diff_plot_path,
                "comparison": cmp_plot_path,
            }
        else:
            plot_files[vector] = {"difference": None, "comparison": None}

    if plt is None:
        add_case("Matplotlib availability", False, "matplotlib is required to generate plots")
    else:
        add_case("Matplotlib availability", True)

    env_txt = (
        f"timestamp={time.strftime('%Y-%m-%d %H:%M:%S')}\n"
        f"project_root={project_root}\n"
        f"node_executable={node_executable}\n"
        f"node_version={node_version}\n"
        f"usig_cli={usig_cli}\n"
        f"golden_raw_dataset={golden_raw_csv}\n"
        f"golden_debug_dataset={golden_debug_csv}\n"
        f"computed_debug_dataset={computed_debug_csv}\n"
        f"tolerance_percent={KPI_TOLERANCE_PERCENT}\n"
        f"scalar_command={' '.join(scalar_command) if scalar_command else 'N/A'}\n"
        f"debug_command={' '.join(cli_command) if cli_command else 'N/A'}\n"
    )
    save_text(output_folder / "environment.txt", env_txt)

    status_df = pd.DataFrame(test_rows)
    kpi_df = pd.DataFrame(kpi_rows)
    stats_df = pd.DataFrame(vector_stats_rows)

    passed, failed, skipped = check_status(report)
    overall_status = "PASS" if failed == 0 else "FAIL"

    document = Document()
    document.add_heading("SINL REGRESSION VALIDATION REPORT", level=0)

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(f"Overall result: {overall_status}")
    document.add_paragraph("Regression checks summary (PASS/FAIL/SKIP):")
    pd_to_docx_table(document, status_df[["Test", "Status", "Reason"]])

    document.add_heading("2. Test Environment", level=1)
    document.add_heading("2.1 Runtime", level=2)
    document.add_paragraph(f"Timestamp: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    document.add_paragraph(f"Node.js version: {node_version}")
    document.add_paragraph(f"USIG CLI path: {usig_cli}")
    document.add_paragraph(f"Project root: {project_root}")

    document.add_heading("2.2 Input Data", level=2)
    document.add_paragraph(f"Golden raw dataset: {golden_raw_csv}")
    document.add_paragraph(f"Golden debug dataset: {golden_debug_csv}")
    document.add_paragraph(f"Computed debug dataset: {computed_debug_csv}")

    document.add_heading("2.3 Regression Configuration", level=2)
    document.add_paragraph(f"Tolerance: {format_num(KPI_TOLERANCE_PERCENT)}%")
    document.add_paragraph(f"Display decimal places: {DISPLAY_DECIMAL_PLACES}")

    document.add_heading("3. Debug Table Generation", level=1)
    document.add_heading("3.1 CLI Command", level=2)
    document.add_paragraph("Scalar KPI command:")
    document.add_paragraph(" ".join(scalar_command) if scalar_command else "N/A")
    document.add_paragraph("Debug table command:")
    document.add_paragraph(" ".join(cli_command) if cli_command else "N/A")

    document.add_heading("3.2 Output Validation", level=2)
    document.add_paragraph(f"CLI execution status: {'PASS' if any(r['check'] == 'CLI debug table execution' and r['passed'] for r in report) else 'FAIL'}")
    document.add_paragraph(f"Debug table generated: {'PASS' if computed_exists else 'FAIL'}")

    document.add_heading("3.3 Structural Validation", level=2)
    golden_rows = len(golden_df.index) if golden_df is not None else "N/A"
    computed_rows = len(computed_df.index) if computed_df is not None else "N/A"
    document.add_paragraph(f"Golden row count: {golden_rows}")
    document.add_paragraph(f"Computed row count: {computed_rows}")
    if golden_df is not None:
        for v in EXPECTED_DEBUG_COLUMNS:
            if v in golden_df.columns:
                document.add_paragraph(f"Golden NaN count ({v}): {format_num(int(golden_df[v].isna().sum()))}")
    if computed_df is not None:
        for v in EXPECTED_DEBUG_COLUMNS:
            if v in computed_df.columns:
                document.add_paragraph(f"Computed NaN count ({v}): {format_num(int(computed_df[v].isna().sum()))}")

    document.add_heading("4. Scalar KPI Regression", level=1)
    pd_to_docx_table(document, kpi_df[["KPI", "Golden", "Computed", "Difference %", "Status"]])

    document.add_heading("5. Statistical Regression", level=1)
    if not stats_df.empty:
        pd_to_docx_table(document, stats_df[["Vector", "Golden Mean", "Computed Mean", "Mean Diff %", "Golden Std", "Computed Std", "Std Diff %", "Golden NaNs", "Computed NaNs", "Mean", "Std", "Overall"]])

    for i, vector in enumerate(EXPECTED_DEBUG_COLUMNS, start=1):
        document.add_heading(f"5.{i} {vector}", level=2)
        row_match = stats_df[stats_df["Vector"] == vector]
        if row_match.empty:
            document.add_paragraph("Overall: SKIP")
            document.add_paragraph("Interpretation: Skipped because structural validation failed.")
            continue
        row = row_match.iloc[0]
        document.add_paragraph(f"Overall: {row['Overall']}")
        document.add_paragraph(f"Golden mean: {row['Golden Mean']}")
        document.add_paragraph(f"Computed mean: {row['Computed Mean']}")
        document.add_paragraph(f"Mean difference: {row['Mean Diff %']} %")
        document.add_paragraph(f"Golden std: {row['Golden Std']}")
        document.add_paragraph(f"Computed std: {row['Computed Std']}")
        document.add_paragraph(f"Std difference: {row['Std Diff %']} %")
        document.add_paragraph(f"Golden NaN count: {format_num(row['Golden NaNs'])}")
        document.add_paragraph(f"Computed NaN count: {format_num(row['Computed NaNs'])}")
        document.add_paragraph(f"NaN count match: {row.get('NaN Count Match', 'N/A')}")
        document.add_paragraph(f"NaN location match: {row.get('NaN Location Match', 'N/A')}")
        document.add_paragraph(f"Mean status: {row['Mean']}")
        document.add_paragraph(f"Std status: {row['Std']}")
        document.add_paragraph(f"Tolerance: {format_num(KPI_TOLERANCE_PERCENT)}%")
        document.add_paragraph(f"Interpretation: {row['Interpretation']}")

    document.add_heading("6. Vector Difference Plots", level=1)
    for i, vector in enumerate(EXPECTED_DEBUG_COLUMNS, start=1):
        document.add_heading(f"6.{i} {vector}", level=2)
        paths = plot_files.get(vector, {})
        diff_plot = paths.get("difference")
        cmp_plot = paths.get("comparison")
        if diff_plot and diff_plot.exists():
            document.add_picture(str(diff_plot), width=Inches(6.8))
            document.add_paragraph(
                f"Figure: {vector} difference between computed and golden SINL debug vectors. "
                f"y = computed - golden; the horizontal zero line represents exact agreement."
            )
        else:
            document.add_paragraph("Difference plot unavailable.")

        if cmp_plot and cmp_plot.exists():
            document.add_picture(str(cmp_plot), width=Inches(6.8))
            document.add_paragraph(
                f"Figure: {vector} comparison of golden and computed SINL debug vectors."
            )

    document.add_heading("7. Final Regression Result", level=1)
    document.add_paragraph(f"Overall regression status: {overall_status}")
    document.add_paragraph(f"Total checks: {len(report)}, Passed: {passed}, Failed: {failed}, Skipped: {skipped}")

    docx_path = output_folder / "sinl_regression_report.docx"
    document.save(docx_path)

    lines = []
    lines.append("=" * 80)
    lines.append("USIG SINL REGRESSION VALIDATION REPORT")
    lines.append("=" * 80)
    for item in report:
        if item.get("skipped", False):
            status = "SKIP"
        else:
            status = "PASS" if item.get("passed", False) else "FAIL"
        reason = f" ({item['reason']})" if item.get("reason") else ""
        lines.append(f"{status:<5} : {item['check']}{reason}")
    lines.append("-" * 80)
    lines.append(f"TOTAL: {len(report)} checks, {passed} passed, {failed} failures, {skipped} skipped")
    lines.append(f"DOCX: {docx_path}")
    lines.append("=" * 80)
    summary = "\n".join(lines)
    print(summary)
    save_text(output_folder / "sinl_regression_report.txt", summary)

    return 0 if failed == 0 else 1


if __name__ == "__main__":
    result = main()
    if result != 0:
        raise SystemExit(result)
